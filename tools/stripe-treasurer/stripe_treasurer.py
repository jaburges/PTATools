#!/usr/bin/env python3
"""
Monthly Stripe report for the Wilder PTSA treasurer, split by product category.

    python3 tools/stripe-treasurer/stripe_treasurer.py 2026-09
    python3 tools/stripe-treasurer/stripe_treasurer.py          # last month

Needs:
  * Stripe CLI logged in to the Wilder account under its own project name:
        stripe login --project-name wilderptsa
    The CLI's live key expires after 90 days; rerun the login when it does.
  * az CLI logged in with access to the Wilder subscription. WooCommerce
    orders are read (read-only) through the wp-db-sql Container Apps job.

Writes reports/<month>/wilder-stripe-<month>.xlsx next to this script:
  Summary        money in and out during the month, by category
  Payouts        each bank deposit that arrived in the month, by category
  Payout detail  the Stripe transactions that make up each deposit
  Transactions   every Stripe transaction created in the month
  Issues         where Stripe and the website disagree

Categories come from WooCommerce product categories, mapped through
categories.json. Each Stripe charge is split across the categories of the
order it paid for, in proportion to the order's line totals; its Stripe
fee is split the same way. Refunds follow the refunded items when the
website refund lists them, otherwise the order's proportions.
"""

from __future__ import annotations

import argparse
import datetime as dt
import html
import json
import subprocess
import sys
import tempfile
import time
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from zoneinfo import ZoneInfo

HERE = Path(__file__).resolve().parent
TZ = ZoneInfo("America/Los_Angeles")

STRIPE_PROJECT = "wilderptsa"
SUBSCRIPTION = "97f6936d-7300-4a49-a2ad-cbfee3b28e00"
RESOURCE_GROUP = "PTSAWebsite"
SQL_JOB = "wp-db-sql"
DB_HOST = "wilderptsa-wpdb-small.mysql.database.azure.com"
DB_USER = "ptsadbadmin"
DB_NAME = "wilderptsa_wp"
LOG_WORKSPACE = "8b7b736f-f468-4fd5-abef-e0df64360a09"

CHARGE_TYPES = ("charge", "payment")
REFUND_TYPES = ("refund", "payment_refund")


# --------------------------------------------------------------------------
# Stripe
# --------------------------------------------------------------------------

def stripe_get(path, params):
    args = ["stripe", "get", path, "--live", "--project-name", STRIPE_PROJECT]
    for key, value in params:
        args += ["-d", f"{key}={value}"]
    out = subprocess.run(args, capture_output=True, text=True)
    try:
        data = json.loads(out.stdout)
    except json.JSONDecodeError:
        sys.exit(f"Stripe CLI failed: {out.stderr.strip() or out.stdout.strip()}\n"
                 f"If the key expired, run: stripe login --project-name {STRIPE_PROJECT}")
    if "error" in data:
        sys.exit(f"Stripe API error on {path}: {data['error'].get('message')}")
    return data


def stripe_list(path, params):
    rows, after = [], None
    while True:
        page = list(params) + [("limit", 100)] + ([("starting_after", after)] if after else [])
        data = stripe_get(path, page)
        rows += data["data"]
        if not data.get("has_more"):
            return rows
        after = data["data"][-1]["id"]


TXN_EXPAND = [("expand[]", "data.source"), ("expand[]", "data.source.payment_intent")]


def order_id_of(txn):
    src = txn.get("source")
    if not isinstance(src, dict):
        return None
    metas = [src.get("metadata") or {}]
    pi = src.get("payment_intent")
    if isinstance(pi, dict):
        metas.append(pi.get("metadata") or {})
    for meta in metas:
        value = meta.get("order_id")
        if value and str(value).isdigit():
            return int(value)
    return None


# --------------------------------------------------------------------------
# WooCommerce, through the read-only SQL job
# --------------------------------------------------------------------------

def az(args):
    out = subprocess.run(["az"] + args, capture_output=True, text=True)
    if out.returncode != 0:
        sys.exit(f"az {' '.join(args[:3])} failed: {out.stderr.strip()}")
    return out.stdout.strip()


def wc_sql(order_ids, gmt_from, gmt_to):
    ids = ",".join(str(i) for i in sorted(order_ids)) or "0"
    clean = "REPLACE(REPLACE(REPLACE({}, '|', '/'), CHAR(10), ' '), CHAR(13), ' ')"
    cats = ("IFNULL((SELECT GROUP_CONCAT(t.name ORDER BY t.name SEPARATOR ';') FROM wp_term_relationships tr "
            "JOIN wp_term_taxonomy tt ON tt.term_taxonomy_id = tr.term_taxonomy_id AND tt.taxonomy = 'product_cat' "
            "JOIN wp_terms t ON t.term_id = tt.term_id WHERE tr.object_id = pid.meta_value), '')")
    items = (
        "SELECT CONCAT_WS('|', '{tag}', oi.order_id, oi.order_item_id, oi.order_item_type, IFNULL(pid.meta_value, ''), "
        + clean.format("oi.order_item_name") + ", "
        "IFNULL(lt.meta_value, IFNULL(cost.meta_value, 0)), " + clean.format(cats) + ") "
        "FROM wp_woocommerce_order_items oi JOIN {table} ON {table}.id = oi.order_id "
        "LEFT JOIN wp_woocommerce_order_itemmeta pid ON pid.order_item_id = oi.order_item_id AND pid.meta_key = '_product_id' "
        "LEFT JOIN wp_woocommerce_order_itemmeta lt ON lt.order_item_id = oi.order_item_id AND lt.meta_key = '_line_total' "
        "LEFT JOIN wp_woocommerce_order_itemmeta cost ON cost.order_item_id = oi.order_item_id AND cost.meta_key = 'cost' "
        "WHERE oi.order_item_type IN ('line_item', 'fee', 'shipping');"
    )
    return " ".join([
        "SET SESSION group_concat_max_len = 10000;",
        "CREATE TEMPORARY TABLE ids (id BIGINT UNSIGNED PRIMARY KEY);",
        f"INSERT IGNORE INTO ids SELECT id FROM wp_wc_orders WHERE type = 'shop_order' AND id IN ({ids});",
        "INSERT IGNORE INTO ids SELECT o.id FROM wp_wc_orders o JOIN wp_wc_order_operational_data od ON od.order_id = o.id "
        f"WHERE o.type = 'shop_order' AND o.payment_method LIKE 'stripe%' AND od.date_paid_gmt >= '{gmt_from}' AND od.date_paid_gmt < '{gmt_to}';",
        "INSERT IGNORE INTO ids SELECT r.parent_order_id FROM wp_wc_orders r JOIN wp_wc_orders p ON p.id = r.parent_order_id "
        f"WHERE r.type = 'shop_order_refund' AND p.payment_method LIKE 'stripe%' AND r.date_created_gmt >= '{gmt_from}' AND r.date_created_gmt < '{gmt_to}';",
        "CREATE TEMPORARY TABLE rids (id BIGINT UNSIGNED PRIMARY KEY);",
        "INSERT INTO rids SELECT r.id FROM wp_wc_orders r JOIN ids ON ids.id = r.parent_order_id WHERE r.type = 'shop_order_refund';",
        "SELECT CONCAT_WS('|', 'O', o.id, o.status, IFNULL(o.payment_method, ''), o.total_amount, IFNULL(o.tax_amount, 0), "
        "IFNULL(od.date_paid_gmt, ''), IFNULL(o.date_created_gmt, '')) "
        "FROM wp_wc_orders o JOIN ids ON ids.id = o.id LEFT JOIN wp_wc_order_operational_data od ON od.order_id = o.id;",
        items.format(tag="L", table="ids"),
        "SELECT CONCAT_WS('|', 'R', r.id, r.parent_order_id, r.total_amount, IFNULL(r.tax_amount, 0), r.date_created_gmt) "
        "FROM wp_wc_orders r JOIN rids ON rids.id = r.id;",
        items.format(tag="RL", table="rids"),
        "SELECT CONCAT_WS('|', 'E', (SELECT COUNT(*) FROM ids), (SELECT COUNT(*) FROM rids));",
    ])


def run_sql_job(sql, execution=None):
    if execution:
        return wait_for_output(execution)
    script = (f'mysql -h {DB_HOST} -u {DB_USER} --ssl-mode=REQUIRED -D {DB_NAME} '
              '-N -B -r -e "$Q"')
    body = {"containers": [{
        "name": SQL_JOB, "image": "docker.io/library/mysql:8",
        "command": ["/bin/sh"], "args": ["-c", script],
        "env": [{"name": "MYSQL_PWD", "secretRef": "db-password"}, {"name": "Q", "value": sql}],
        "resources": {"cpu": 0.25, "memory": "0.5Gi"},
    }]}
    with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False) as fh:
        json.dump(body, fh)
    url = (f"https://management.azure.com/subscriptions/{SUBSCRIPTION}/resourceGroups/{RESOURCE_GROUP}"
           f"/providers/Microsoft.App/jobs/{SQL_JOB}/start?api-version=2024-03-01")
    execution = az(["rest", "--method", "post", "--url", url, "--body", f"@{fh.name}", "--query", "name", "-o", "tsv"])
    Path(fh.name).unlink()
    print(f"  website query running as {execution} …", flush=True)

    for _ in range(60):
        status = az(["containerapp", "job", "execution", "show", "-g", RESOURCE_GROUP, "-n", SQL_JOB,
                     "--job-execution-name", execution, "--subscription", SUBSCRIPTION,
                     "--query", "properties.status", "-o", "tsv"])
        if status in ("Succeeded", "Failed"):
            break
        time.sleep(10)
    if status != "Succeeded":
        sys.exit(f"website query {execution} ended {status}")
    return wait_for_output(execution)


def wait_for_output(execution):
    # Console output reaches Log Analytics a minute or two after the job ends.
    query = (f"ContainerAppConsoleLogs_CL | where ContainerGroupName_s startswith '{execution}' "
             "| project Log_s")
    last = -1
    for _ in range(40):
        lines = log_query(query)
        if any(line.startswith("E|") for line in lines) and len(lines) == last:
            return lines
        last = len(lines)
        time.sleep(15)
    sys.exit(f"website query {execution} output never arrived in Log Analytics; "
             f"rerun later with --execution {execution}")


def log_query(query):
    # az's own log-analytics command signs in as the default subscription's
    # account, which is not the Wilder one when another PTA is the default.
    import urllib.request
    token = az(["account", "get-access-token", "--subscription", SUBSCRIPTION,
                "--resource", "https://api.loganalytics.io", "--query", "accessToken", "-o", "tsv"])
    req = urllib.request.Request(
        f"https://api.loganalytics.io/v1/workspaces/{LOG_WORKSPACE}/query",
        data=json.dumps({"query": query, "timespan": "P1D"}).encode(),
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=60) as resp:
        table = json.load(resp)["tables"][0]
    col = [c["name"] for c in table["columns"]].index("Log_s")
    return [row[col].rstrip("\n") for row in table["rows"]]


def cents(value):
    return int((Decimal(str(value or 0)) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def parse_wc(lines):
    orders, refunds = {}, {}
    seen_items = set()
    for line in lines:
        parts = line.split("|")
        tag = parts[0]
        if tag == "O" and len(parts) >= 8:
            oid = int(parts[1])
            orders[oid] = {"status": parts[2], "method": parts[3], "total": cents(parts[4]),
                           "tax": cents(parts[5]), "paid_gmt": parts[6], "created_gmt": parts[7],
                           "items": [], "refunds": []}
        elif tag == "R" and len(parts) >= 6:
            refunds[int(parts[1])] = {"order": int(parts[2]), "total": cents(parts[3]),
                                      "tax": cents(parts[4]), "created_gmt": parts[5], "items": []}
        elif tag == "E":
            expected = (int(parts[1]), int(parts[2]))
    for line in lines:
        parts = line.split("|")
        if parts[0] not in ("L", "RL") or len(parts) < 8:
            continue
        key = (parts[0], parts[2])
        if key in seen_items:
            continue
        seen_items.add(key)
        item = {"type": parts[3], "product": parts[4], "name": html.unescape(parts[5]),
                "amount": cents(parts[6]), "cats": [html.unescape(c) for c in parts[7].split(";") if c]}
        owner = orders if parts[0] == "L" else refunds
        if int(parts[1]) in owner:
            owner[int(parts[1])]["items"].append(item)
    for rid, ref in refunds.items():
        if ref["order"] in orders:
            ref["id"] = rid
            orders[ref["order"]]["refunds"].append(ref)
    if (len(orders), len(refunds)) != expected:
        sys.exit(f"website query returned {len(orders)} orders / {len(refunds)} refunds, expected {expected}; rerun")
    return orders


# --------------------------------------------------------------------------
# Allocation
# --------------------------------------------------------------------------

class Categories:
    def __init__(self, config):
        self.cfg = config
        self.priority = config["priority"]
        self.aliases = config.get("aliases", {})

    def label_for(self, item):
        if item["type"] == "fee":
            return self.cfg["fee_lines"]
        if item["type"] == "shipping":
            return self.cfg["shipping"]
        names = sorted({self.aliases.get(c, c) for c in item["cats"]} - {"Uncategorized"})
        if not names:
            return self.cfg["uncategorized"]
        ranked = sorted(names, key=lambda n: (self.priority.index(n) if n in self.priority else len(self.priority), n))
        return ranked[0]

    def order(self, labels):
        fixed = [self.cfg[k] for k in ("fee_lines", "shipping", "tax", "uncategorized", "order_missing", "no_order")]
        def key(label):
            if label in self.priority:
                return (0, self.priority.index(label), label)
            if label.startswith("Stripe: "):
                return (3, 0, label)
            if label in fixed:
                return (2, fixed.index(label), label)
            return (1, 0, label)
        return sorted(labels, key=key)


def split(amount, weights):
    """Split integer cents across labels in proportion to weights, exactly."""
    weights = {k: abs(v) for k, v in weights.items() if v}
    if not amount:
        return {}
    if not weights:
        return None
    total = sum(weights.values())
    sign = -1 if amount < 0 else 1
    amount = abs(amount)
    shares = {k: amount * w // total for k, w in weights.items()}
    rest = amount - sum(shares.values())
    for k in sorted(weights, key=lambda k: (amount * weights[k] % total), reverse=True)[:rest]:
        shares[k] += 1
    return {k: sign * v for k, v in shares.items() if v}


def order_weights(order, cats):
    weights = defaultdict(int)
    for item in order["items"]:
        weights[cats.label_for(item)] += item["amount"]
    weights[cats.cfg["tax"]] += order["tax"]
    return weights


class Allocator:
    def __init__(self, orders, cats):
        self.orders = orders
        self.cats = cats
        self.issues = []
        self.used_refunds = set()
        self.cache = {}

    def issue(self, kind, detail, ref="", amount=None):
        self.issues.append({"kind": kind, "ref": ref, "detail": detail, "amount": amount})

    def allocate(self, txn):
        if txn["id"] in self.cache:
            return self.cache[txn["id"]]
        t = txn["type"]
        oid = order_id_of(txn)
        note = ""
        order = self.orders.get(oid) if oid else None
        if t in CHARGE_TYPES or t in REFUND_TYPES:
            if oid is None:
                weights = {self.cats.cfg["no_order"]: 1}
                desc = (txn.get("source") or {}).get("description") or txn.get("description") or ""
                note = f"No website order on this payment. {desc}".strip()
                if t in CHARGE_TYPES:
                    self.issue("Stripe payment with no website order", note, txn["id"], txn["amount"])
            elif order is None:
                weights = {self.cats.cfg["order_missing"]: 1}
                note = f"Order {oid} is not on the website"
                self.issue("Order not found on website", note, txn["id"], txn["amount"])
            elif t in CHARGE_TYPES:
                weights = order_weights(order, self.cats)
                if txn["amount"] != order["total"]:
                    note = f"Stripe {txn['amount'] / 100:.2f} vs order total {order['total'] / 100:.2f}"
                    self.issue("Payment differs from order total", f"Order {oid}: {note}", txn["id"], txn["amount"])
            else:
                weights, note = self.refund_weights(order, oid, txn)
        else:
            weights = {f"Stripe: {t.replace('_', ' ')}": 1}
            note = txn.get("description") or ""

        gross = split(txn["amount"], weights) or {}
        fee = split(-txn["fee"], weights) or {}
        result = {"order": oid, "gross": gross, "fee": fee, "note": note}
        self.cache[txn["id"]] = result
        return result

    def refund_weights(self, order, oid, txn):
        amount = -txn["amount"]
        for ref in order["refunds"]:
            if ref["id"] in self.used_refunds or -ref["total"] != amount:
                continue
            self.used_refunds.add(ref["id"])
            weights = defaultdict(int)
            for item in ref["items"]:
                weights[self.cats.label_for(item)] += -item["amount"]
            weights[self.cats.cfg["tax"]] += -ref["tax"]
            if sum(weights.values()) == amount:
                return weights, f"Website refund #{ref['id']} by item"
            return order_weights(order, self.cats), f"Website refund #{ref['id']} has no item detail; split like the order"
        self.issue("Stripe refund with no matching website refund",
                   f"Order {oid}: refund of {amount / 100:.2f}", txn["id"], txn["amount"])
        return order_weights(order, self.cats), "No matching website refund; split like the order"


# --------------------------------------------------------------------------
# Report
# --------------------------------------------------------------------------

def local_date(ts):
    return dt.datetime.fromtimestamp(ts, TZ).strftime("%Y-%m-%d")


def fmt_split(parts):
    return "; ".join(f"{k} {v / 100:,.2f}" for k, v in sorted(parts.items(), key=lambda kv: -abs(kv[1])))


def build(month, execution=None):
    year, mon = (int(x) for x in month.split("-"))
    start = dt.datetime(year, mon, 1, tzinfo=TZ)
    end = dt.datetime(year + (mon == 12), mon % 12 + 1, 1, tzinfo=TZ)
    utc = dt.timezone.utc
    gmt_from = start.astimezone(utc).strftime("%Y-%m-%d %H:%M:%S")
    gmt_to = end.astimezone(utc).strftime("%Y-%m-%d %H:%M:%S")

    print(f"Stripe: transactions created {start:%b %-d} to {end - dt.timedelta(days=1):%b %-d, %Y} …", flush=True)
    month_txns = stripe_list("/v1/balance_transactions",
                             [("created[gte]", int(start.timestamp())), ("created[lt]", int(end.timestamp()))] + TXN_EXPAND)
    # arrival_date is midnight UTC of the bank date.
    a_from = int(dt.datetime(year, mon, 1, tzinfo=utc).timestamp())
    a_to = int(dt.datetime(year + (mon == 12), mon % 12 + 1, 1, tzinfo=utc).timestamp())
    payouts = stripe_list("/v1/payouts", [("arrival_date[gte]", a_from), ("arrival_date[lt]", a_to)])
    payout_txns = {}
    for po in payouts:
        rows = stripe_list("/v1/balance_transactions", [("payout", po["id"])] + TXN_EXPAND)
        payout_txns[po["id"]] = [r for r in rows if r["type"] != "payout"]
    print(f"  {len(month_txns)} transactions, {len(payouts)} payouts", flush=True)

    order_ids = {order_id_of(t) for t in month_txns}
    for rows in payout_txns.values():
        order_ids |= {order_id_of(t) for t in rows}
    order_ids.discard(None)

    print("Website: orders, line items and refunds …", flush=True)
    orders = parse_wc(run_sql_job(wc_sql(order_ids, gmt_from, gmt_to), execution))
    print(f"  {len(orders)} orders", flush=True)

    cats = Categories(json.loads((HERE / "categories.json").read_text()))
    alloc = Allocator(orders, cats)

    for oid, order in orders.items():
        lines = sum(i["amount"] for i in order["items"]) + order["tax"]
        if lines != order["total"]:
            alloc.issue("Order lines don't add up to the order total",
                        f"Order {oid}: lines {lines / 100:.2f}, total {order['total'] / 100:.2f}", str(oid), order["total"])

    summary = defaultdict(lambda: defaultdict(int))
    transactions = []
    for t in sorted(month_txns, key=lambda t: t["created"]):
        if t["type"] == "payout":
            continue
        a = alloc.allocate(t)
        column = "sales" if t["type"] in CHARGE_TYPES else "refunds" if t["type"] in REFUND_TYPES else "other"
        for label, v in a["gross"].items():
            summary[label][column] += v
        for label, v in a["fee"].items():
            summary[label]["fees"] += v
        transactions.append([local_date(t["created"]), t["type"], t["id"], a["order"] or "",
                             t["amount"] / 100, -t["fee"] / 100, t["net"] / 100, fmt_split(a["gross"]), a["note"]])

    charged_orders = set()
    for rows in [month_txns] + list(payout_txns.values()):
        charged_orders |= {order_id_of(t) for t in rows if t["type"] in CHARGE_TYPES}
    for oid, order in sorted(orders.items()):
        if order["method"].startswith("stripe") and gmt_from <= order["paid_gmt"] < gmt_to and oid not in charged_orders:
            alloc.issue("Website order paid by card with no Stripe payment this month",
                        f"Order {oid} ({order['status']}) paid {order['paid_gmt']} UTC", str(oid), order["total"])
        for ref in order["refunds"]:
            if gmt_from <= ref["created_gmt"] < gmt_to and ref["id"] not in alloc.used_refunds and ref["total"]:
                alloc.issue("Website refund with no Stripe refund",
                            f"Order {oid}: refund #{ref['id']} of {-ref['total'] / 100:.2f} on {ref['created_gmt']} UTC",
                            str(oid), ref["total"])

    payout_rows, payout_detail = [], []
    for po in sorted(payouts, key=lambda p: p["arrival_date"]):
        net = defaultdict(int)
        for t in sorted(payout_txns[po["id"]], key=lambda t: t["created"]):
            a = alloc.allocate(t)
            row_net = defaultdict(int)
            for label, v in a["gross"].items():
                row_net[label] += v
            for label, v in a["fee"].items():
                row_net[label] += v
            for label, v in row_net.items():
                net[label] += v
            payout_detail.append([po["id"], dt.datetime.fromtimestamp(po["arrival_date"], utc).strftime("%Y-%m-%d"),
                                  local_date(t["created"]), t["type"], t["id"], a["order"] or "",
                                  t["amount"] / 100, -t["fee"] / 100, t["net"] / 100, fmt_split(row_net)])
        diff = sum(net.values()) - po["amount"]
        if diff:
            alloc.issue("Payout doesn't equal the transactions in it",
                        f"{po['id']}: off by {diff / 100:.2f}", po["id"], po["amount"])
        payout_rows.append({"po": po, "net": net})

    return {"month": month, "start": start, "end": end, "summary": summary, "transactions": transactions,
            "payouts": payout_rows, "payout_detail": payout_detail, "issues": alloc.issues, "cats": cats}


def write_xlsx(report, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    money = '#,##0.00;[Red]-#,##0.00'
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDE6F0")
    cats = report["cats"]

    def header(ws, row, values):
        for col, v in enumerate(values, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.font = bold
            c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="top")

    def widths(ws, cols):
        for i, w in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    month_name = report["start"].strftime("%B %Y")
    ws["A1"] = f"Wilder PTSA Stripe report, {month_name}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("Money in and out of Stripe for transactions dated in the month (Pacific time), by category. "
                "Stripe fees are split across categories in proportion to each payment.")
    header(ws, 4, ["Category", "Sales", "Refunds", "Stripe fees", "Other", "Net"])
    r = 5
    totals = defaultdict(int)
    for label in cats.order(report["summary"].keys()):
        s = report["summary"][label]
        net = s["sales"] + s["refunds"] + s["fees"] + s["other"]
        ws.append([label, s["sales"] / 100, s["refunds"] / 100, s["fees"] / 100, s["other"] / 100, net / 100])
        for k in ("sales", "refunds", "fees", "other"):
            totals[k] += s[k]
        totals["net"] += net
        r += 1
    ws.append(["Total", totals["sales"] / 100, totals["refunds"] / 100, totals["fees"] / 100,
               totals["other"] / 100, totals["net"] / 100])
    for c in ws[r]:
        c.font = bold
    for row in ws.iter_rows(min_row=5, max_row=r, min_col=2, max_col=6):
        for c in row:
            c.number_format = money

    r += 2
    ws.cell(row=r, column=1, value="Paid out to the bank (deposits arriving this month)").font = bold
    paid = sum(p["po"]["amount"] for p in report["payouts"])
    ws.cell(row=r + 1, column=1, value="Payouts")
    ws.cell(row=r + 1, column=2, value=paid / 100).number_format = money
    ws.cell(row=r + 2, column=1, value="Count")
    ws.cell(row=r + 2, column=2, value=len(report["payouts"]))
    ws.cell(row=r + 3, column=1, value="Issues to review")
    ws.cell(row=r + 3, column=2, value=len(report["issues"]))
    widths(ws, [34, 14, 14, 14, 12, 14])
    ws.freeze_panes = "A5"

    labels = cats.order({k for p in report["payouts"] for k in p["net"]})
    ws = wb.create_sheet("Payouts")
    ws["A1"] = "Each Stripe payout that arrived in the bank this month, split by category (net of refunds and fees)."
    header(ws, 3, ["Arrival date", "Payout", "Status", "Amount"] + labels)
    for p in report["payouts"]:
        po = p["po"]
        ws.append([dt.datetime.fromtimestamp(po["arrival_date"], dt.timezone.utc).strftime("%Y-%m-%d"), po["id"],
                   po["status"], po["amount"] / 100] + [p["net"].get(k, 0) / 100 for k in labels])
    last = ws.max_row
    ws.append(["Total", "", "", sum(p["po"]["amount"] for p in report["payouts"]) / 100]
              + [sum(p["net"].get(k, 0) for p in report["payouts"]) / 100 for k in labels])
    for c in ws[ws.max_row]:
        c.font = bold
    for row in ws.iter_rows(min_row=4, max_row=last + 1, min_col=4, max_col=4 + len(labels)):
        for c in row:
            c.number_format = money
    widths(ws, [12, 32, 10, 13] + [14] * len(labels))
    ws.freeze_panes = "E4"

    def table(name, intro, cols, rows, money_cols, col_widths):
        sh = wb.create_sheet(name)
        sh["A1"] = intro
        header(sh, 3, cols)
        for row in rows:
            sh.append(row)
        for row in sh.iter_rows(min_row=4, max_row=sh.max_row):
            for idx in money_cols:
                row[idx].number_format = money
        widths(sh, col_widths)
        sh.freeze_panes = "A4"
        sh.auto_filter.ref = f"A3:{get_column_letter(len(cols))}{max(sh.max_row, 3)}"

    table("Payout detail", "The Stripe transactions included in each payout above.",
          ["Payout", "Arrival date", "Date", "Type", "Stripe ID", "Order", "Amount", "Stripe fee", "Net", "Net by category"],
          report["payout_detail"], (6, 7, 8), [30, 12, 11, 10, 30, 8, 11, 11, 11, 60])
    table("Transactions", "Every Stripe transaction created this month, except payouts.",
          ["Date", "Type", "Stripe ID", "Order", "Amount", "Stripe fee", "Net", "Amount by category", "Note"],
          report["transactions"], (4, 5, 6), [11, 10, 30, 8, 11, 11, 11, 60, 50])
    table("Issues", "Where Stripe and the website disagree. Each one needs a look before the numbers are final.",
          ["Issue", "Reference", "Detail", "Amount"],
          [[i["kind"], i["ref"], i["detail"], (i["amount"] or 0) / 100] for i in report["issues"]],
          (3,), [44, 30, 70, 12])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def print_summary(report):
    cats = report["cats"]
    print(f"\n{report['start']:%B %Y}")
    print(f"{'Category':<30}{'Sales':>12}{'Refunds':>12}{'Fees':>11}{'Net':>12}")
    total = defaultdict(int)
    for label in cats.order(report["summary"].keys()):
        s = report["summary"][label]
        net = s["sales"] + s["refunds"] + s["fees"] + s["other"]
        print(f"{label:<30}{s['sales'] / 100:>12,.2f}{s['refunds'] / 100:>12,.2f}{s['fees'] / 100:>11,.2f}{net / 100:>12,.2f}")
        for k in ("sales", "refunds", "fees"):
            total[k] += s[k]
        total["net"] += net
    print(f"{'Total':<30}{total['sales'] / 100:>12,.2f}{total['refunds'] / 100:>12,.2f}{total['fees'] / 100:>11,.2f}{total['net'] / 100:>12,.2f}")
    print(f"\nPayouts arriving this month: {len(report['payouts'])}, "
          f"{sum(p['po']['amount'] for p in report['payouts']) / 100:,.2f}")
    kinds = defaultdict(int)
    for i in report["issues"]:
        kinds[i["kind"]] += 1
    print(f"Issues to review: {len(report['issues'])}")
    for k, n in kinds.items():
        print(f"  {n:>3}  {k}")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    today = dt.datetime.now(TZ)
    last_month = (today.replace(day=1) - dt.timedelta(days=1)).strftime("%Y-%m")
    parser.add_argument("month", nargs="?", default=last_month, help="YYYY-MM (default: last month)")
    parser.add_argument("--out", type=Path, help="output .xlsx path")
    parser.add_argument("--execution", help="reuse the website query output of an earlier run (same month, within a day)")
    args = parser.parse_args()

    report = build(args.month, args.execution)
    path = args.out or HERE / "reports" / args.month / f"wilder-stripe-{args.month}.xlsx"
    write_xlsx(report, path)
    print_summary(report)
    print(f"\nSaved {path}")


if __name__ == "__main__":
    main()
